import io, re
from datetime import datetime, timedelta
import pandas as pd
import streamlit as st

st.set_page_config(page_title="Work Time Analyzer", page_icon="⏱️", layout="wide")

st.markdown("""
<style>
.stApp{background:linear-gradient(135deg,#07111f,#0d1b2a 50%,#132238);color:#f4f7fb}
.block-container{max-width:1400px;padding-top:1.4rem}
.hero{padding:28px 30px;border-radius:24px;margin-bottom:20px;background:linear-gradient(120deg,#17365d,#0b6b6b,#274e6f);box-shadow:0 18px 50px rgba(0,0,0,.28);animation:fade .7s}
.card{padding:18px;border-radius:18px;background:rgba(255,255,255,.07);border:1px solid rgba(255,255,255,.1)}
@keyframes fade{from{opacity:0;transform:translateY(10px)}to{opacity:1;transform:translateY(0)}}
</style>
<div class="hero"><h1>⏱️ Work Time Analyzer</h1>
<p>Daily & monthly attendance • 09:30–18:30 • WOP-aware overtime</p></div>
""", unsafe_allow_html=True)

USERS={"admin":("Admin@123","admin"),"user1":("Work@123","user"),"user2":("Work@123","user"),"user3":("Work@123","user"),"user4":("Work@123","user")}
if "auth" not in st.session_state: st.session_state.auth=False
if not st.session_state.auth:
    with st.form("login"):
        st.subheader("🔐 Sign in")
        u=st.text_input("Username"); p=st.text_input("Password",type="password")
        if st.form_submit_button("Sign in",use_container_width=True):
            if u in USERS and USERS[u][0]==p:
                st.session_state.auth=True; st.session_state.user=u; st.rerun()
            else: st.error("Invalid username or password.")
    st.stop()

with st.sidebar:
    st.success("Signed in: "+st.session_state.user)
    if st.button("Sign out"): st.session_state.clear(); st.rerun()
    st.markdown("### Company rules")
    st.write("Start: **09:30**")
    st.write("End: **18:30**")
    st.write("Saturday: **WOP working day**")
    st.write("Sunday: **WO weekly off**")

def hm(n):
    n=max(0,int(round(n)))
    return f"{n//60:02d}:{n%60:02d}"

def parse_hm(v):
    if v is None or pd.isna(v): return 0
    if isinstance(v,timedelta): return int(v.total_seconds()//60)
    s=str(v).strip()
    if not s or s.lower() in {"nan","nat","none"}: return 0
    m=re.match(r"^(\d+):(\d{1,2})(?::(\d{1,2}))?$",s)
    if m:return int(m.group(1))*60+int(m.group(2))
    try:
        x=float(s)
        return int(x*1440) if 0<=x<1 else int(x)
    except:return 0

def parse_clock(v):
    if v is None or pd.isna(v): return None
    if isinstance(v,datetime): return v
    s=str(v).strip()
    if not s or s.lower() in {"nan","nat","none"}: return None
    for fmt in ("%H:%M:%S","%H:%M","%I:%M %p","%I:%M:%S %p"):
        try:return datetime.strptime(s,fmt)
        except:pass
    return None

def actual_duration(iv,ov):
    a,b=parse_clock(iv),parse_clock(ov)
    if not a or not b:return 0
    d=b-a
    if d.total_seconds()<0:d+=timedelta(days=1)
    return int(d.total_seconds()//60)

def norm_status(v):
    s=str(v).strip().upper()
    if s in {"P","PRESENT"}:return "P"
    if s in {"A","ABSENT"}:return "A"
    if s in {"WOP","SAT","SATURDAY"}:return "WOP"
    if s in {"WO","SUN","SUNDAY"}:return "WO"
    return s

def report_period(raw):
    for i in range(min(15,len(raw))):
        for v in raw.iloc[i].tolist():
            if pd.notna(v):
                m=re.search(r"([A-Z][a-z]{2}\s+\d{2}\s+\d{4}).*?([A-Z][a-z]{2}\s+\d{2}\s+\d{4})",str(v))
                if m:
                    try:return pd.to_datetime(m.group(1)),pd.to_datetime(m.group(2))
                    except:pass
    return None,None

def parse_detailed_report(file_bytes):
    """Parser for the actual Monthly Status Report (Detailed Work Duration)
    format used by sep2026.xlsx: each employee is a vertical block with
    Status/InTime/OutTime/Duration/Late By/Early By/OT/Shift rows and day
    columns such as 1 T, 2 W, 5 St, 6 S.
    """
    raw=pd.read_excel(io.BytesIO(file_bytes),sheet_name=0,header=None)
    start,end=report_period(raw)
    if start is None:
        start=pd.Timestamp(datetime.now().year,datetime.now().month,1)
    day_cols=[]
    if len(raw)>6:
        for c,v in enumerate(raw.iloc[6].tolist()):
            if pd.notna(v):
                m=re.match(r"^\s*(\d{1,2})\s+",str(v))
                if m:
                    day=int(m.group(1))
                    if 1<=day<=31: day_cols.append((c,day))
    if not day_cols:
        raise ValueError("Could not find the daily columns in row 7 of the report.")
    employees=[]
    for r in range(len(raw)):
        first=str(raw.iloc[r,0]).strip() if pd.notna(raw.iloc[r,0]) else ""
        if first.lower()!="employee:": continue
        emp_cell=raw.iloc[r,3] if raw.shape[1]>3 else ""
        emp=str(emp_cell).strip() if pd.notna(emp_cell) else ""
        # Expected form "1 : VANI"; keep the name after colon.
        if ":" in emp: emp=emp.split(":",1)[1].strip()
        if not emp or emp.lower()=="nan": continue
        # Find the labelled rows belonging to this employee.
        labels={}
        rr=r+1
        while rr<len(raw) and rr<r+15:
            label=str(raw.iloc[rr,0]).strip().lower() if pd.notna(raw.iloc[rr,0]) else ""
            if label: labels[label]=rr
            if label=="shift": break
            rr+=1
        status_r=labels.get("status")
        in_r=labels.get("intime")
        out_r=labels.get("outtime")
        dur_r=labels.get("duration")
        late_r=labels.get("late by")
        early_r=labels.get("early by")
        ot_r=labels.get("ot")
        for c,day in day_cols:
            status=norm_status(raw.iloc[status_r,c]) if status_r is not None and c<raw.shape[1] else ""
            iv=raw.iloc[in_r,c] if in_r is not None else ""
            ov=raw.iloc[out_r,c] if out_r is not None else ""
            dur_source=parse_hm(raw.iloc[dur_r,c]) if dur_r is not None else 0
            actual=actual_duration(iv,ov)
            # WOP duration is often 00:00 in the source. Use In/Out for WOP.
            work=actual if status=="WOP" and actual>0 else dur_source
            # For other P days, source duration is authoritative when present.
            if work==0 and actual>0 and status=="P": work=actual
            source_ot=parse_hm(raw.iloc[ot_r,c]) if ot_r is not None else 0
            # Company rule: Saturday WOP is a working day, but WOP OT is excluded
            # from regular Total OT. Preserve it separately for transparency.
            regular_ot=0 if status=="WOP" else source_ot
            late=parse_hm(raw.iloc[late_r,c]) if late_r is not None else 0
            early=parse_hm(raw.iloc[early_r,c]) if early_r is not None else 0
            if not late and status=="P":
                t=parse_clock(iv)
                if t:
                    late=max(0,int((t-datetime(1900,1,1,9,30)).total_seconds()//60))
            if not early and status=="P":
                t=parse_clock(ov)
                if t:
                    early=max(0,int((datetime(1900,1,1,18,30)-t).total_seconds()//60))
            # Construct real date from report month + day.
            try:
                dt=pd.Timestamp(start.year,start.month,day)
                if end is not None and dt>end: continue
            except: continue
            employees.append({
                "Date":dt,"Employee":emp,"Status":status,
                "In Time":str(iv) if pd.notna(iv) else "",
                "Out Time":str(ov) if pd.notna(ov) else "",
                "Work Duration":work,"Regular OT":regular_ot,
                "Late By":late,"Early By":early,
                "WOP OT Excluded":source_ot if status=="WOP" else 0
            })
    if not employees:
        raise ValueError("No employee attendance blocks were found. Expected rows labelled Employee:, Status, InTime, OutTime, Duration, OT.")
    return pd.DataFrame(employees)

def summary(df):
    out=[]
    for name,g in df.groupby("Employee",sort=True):
        p=int((g.Status=="P").sum()); a=int((g.Status=="A").sum())
        wop=int((g.Status=="WOP").sum()); wo=int((g.Status=="WO").sum())
        work=int(g["Work Duration"].sum()); ot=int(g["Regular OT"].sum())
        late=int(g["Late By"].sum()); early=int(g["Early By"].sum())
        out.append({
            "Employee":name,"Total Work Duration":hm(work),"Total OT":hm(ot),
            "Present":p,"Absent":a,"WOP":wop,"WeeklyOff":wo,
            "Holidays":0,"Leaves Taken":0,"Late By Hrs":hm(late),
            "Late By Days":int((g["Late By"]>0).sum()),"Early By Hrs":hm(early),
            "Early Going By Days":int((g["Early By"]>0).sum()),
            "WOP OT Excluded":hm(int(g["WOP OT Excluded"].sum())),
            "Total Duration (+OT)":hm(work+ot),
            "Average Working Hrs":hm(round(work/p) if p else 0)
        })
    return pd.DataFrame(out)

st.markdown("### 📁 Upload attendance report(s)")
files=st.file_uploader("Upload the company's detailed attendance Excel report",type=["xlsx","xls"],accept_multiple_files=True)

if not files:
    st.info("Upload your attendance Excel file to begin.")
    st.stop()

frames=[]; errors=[]
for f in files:
    try:
        d=parse_detailed_report(f.getvalue())
        d["Source File"]=f.name
        frames.append(d)
    except Exception as e:
        errors.append(f"{f.name}: {e}")

if errors:
    for e in errors: st.warning(e)
if not frames:
    st.error("No employee attendance rows could be detected. This version expects the Monthly Status Report (Detailed Work Duration) format.")
    st.stop()

data=pd.concat(frames,ignore_index=True).drop_duplicates(
    subset=["Date","Employee","Status","In Time","Out Time"],keep="first"
)

st.success(f"Loaded {len(files)} file(s) • {data.Employee.nunique()} employees • {data.Date.nunique()} dates")

t1,t2,t3,t4=st.tabs(["📊 Dashboard","📅 Daily / Monthly","👤 Employee","📤 Export"])
with t1:
    s=summary(data)
    c=st.columns(4)
    c[0].metric("Employees",data.Employee.nunique())
    c[1].metric("Present records",int((data.Status=="P").sum()))
    c[2].metric("Regular OT",hm(data["Regular OT"].sum()))
    c[3].metric("WOP OT excluded",hm(data["WOP OT Excluded"].sum()))
    st.dataframe(s,use_container_width=True,hide_index=True)

with t2:
    view=data.copy(); view["Date"]=view["Date"].dt.strftime("%Y-%m-%d")
    st.dataframe(view,use_container_width=True,hide_index=True)
    st.download_button("⬇️ Download CSV",view.to_csv(index=False).encode(),"work_time_daily.csv","text/csv")

with t3:
    name=st.selectbox("Select employee",sorted(data.Employee.unique()))
    e=data[data.Employee==name].copy()
    st.dataframe(summary(e),use_container_width=True,hide_index=True)
    e["Date"]=e["Date"].dt.strftime("%Y-%m-%d")
    st.dataframe(e,use_container_width=True,hide_index=True)

with t4:
    period=st.radio("Report period",["Daily","Monthly"],horizontal=True)
    who=st.radio("Employees",["All Employees","Selected Employees"],horizontal=True)
    if who=="Selected Employees":
        selected=st.multiselect("Select employees",sorted(data.Employee.unique()))
        ex=data[data.Employee.isin(selected)].copy()
    else: ex=data.copy()
    if period=="Daily":
        ds=sorted(data.Date.dt.strftime("%Y-%m-%d").unique())
        day=st.selectbox("Select date",ds)
        ex=ex[ex.Date.dt.strftime("%Y-%m-%d")==day]
    else:
        ms=sorted(data.Date.dt.to_period("M").astype(str).unique())
        month=st.selectbox("Select month",ms)
        ex=ex[ex.Date.dt.to_period("M").astype(str)==month]
    st.dataframe(summary(ex),use_container_width=True,hide_index=True)
    fmt=st.radio("Format",["PDF","Word","Image"],horizontal=True)
    s=summary(ex)
    if fmt=="PDF":
        from reportlab.platypus import SimpleDocTemplate,Table,TableStyle,Paragraph
        from reportlab.lib.pagesizes import landscape,A4
        from reportlab.lib import colors
        from reportlab.lib.styles import getSampleStyleSheet
        b=io.BytesIO(); doc=SimpleDocTemplate(b,pagesize=landscape(A4))
        vals=[list(s.columns)]+s.astype(str).values.tolist()
        tb=Table(vals,repeatRows=1)
        tb.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#17365d")),("TEXTCOLOR",(0,0),(-1,0),colors.white),("GRID",(0,0),(-1,-1),.25,colors.grey),("FONTSIZE",(0,0),(-1,-1),6)]))
        doc.build([Paragraph("Work Time Analyzer Report",getSampleStyleSheet()["Title"]),tb])
        st.download_button("⬇️ Download PDF",b.getvalue(),"work_time_report.pdf","application/pdf")
    elif fmt=="Word":
        from docx import Document
        b=io.BytesIO(); doc=Document(); doc.add_heading("Work Time Analyzer Report",0)
        tb=doc.add_table(rows=1,cols=len(s.columns))
        for i,c in enumerate(s.columns): tb.rows[0].cells[i].text=str(c)
        for _,r in s.iterrows():
            cells=tb.add_row().cells
            for i,c in enumerate(s.columns): cells[i].text=str(r[c])
        doc.save(b)
        st.download_button("⬇️ Download Word",b.getvalue(),"work_time_report.docx","application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    else:
        import matplotlib.pyplot as plt
        b=io.BytesIO(); fig,ax=plt.subplots(figsize=(18,max(4,len(s)*.45))); ax.axis("off")
        tb=ax.table(cellText=s.astype(str).values,colLabels=s.columns,loc="center")
        tb.auto_set_font_size(False); tb.set_fontsize(7); tb.scale(1,1.4)
        fig.savefig(b,format="png",dpi=180,bbox_inches="tight"); plt.close(fig)
        st.download_button("⬇️ Download Image",b.getvalue(),"work_time_report.png","image/png")

st.caption("P = Present • A = Absent • WOP = Saturday working day • WO = Sunday weekly off • WOP OT is excluded from regular Total OT.")
